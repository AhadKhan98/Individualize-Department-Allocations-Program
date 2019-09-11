import pandas
from docx import Document
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

def create_doc(department,account_num,s_5hr,s_10hr,p_10hr,p_12hr,p_15hr,total_primary,thanksgiving_hours,spring_hours,christmas_hours,summer_hours,i):
    document = Document()

    # Table Layout
    table = document.add_table(rows = 4, cols = 12)
    table.style = 'TableGrid'

    # Merging Selected Cells
    table.cell(0,0).merge(table.cell(0,11))
    table.cell(1,0).merge(table.cell(1,1))
    table.cell(1,2).merge(table.cell(1,3))
    table.cell(1,4).merge(table.cell(1,6))
    table.cell(1,7).merge(table.cell(1,11))
    table.cell(2,0).merge(table.cell(2,1))
    table.cell(3,0).merge(table.cell(3,1))

    # Setting Table Labels
    table.cell(1,3).text = "Secondary"
    table.cell(1,4).text = "Primary"
    table.cell(2,2).text = "5hr"
    table.cell(2,3).text = "10hr"
    table.cell(2,4).text = "10hr"
    table.cell(2,5).text = "12hr"
    table.cell(2,6).text = "15hr"
    table.cell(2,7).text = "Total Primary"
    table.cell(2,8).text = "Thanksgiving Hours"
    table.cell(2,9).text = "Spring Hours"
    table.cell(2,10).text = "Christmas Hours"
    table.cell(2,11).text = "Summer Hours"
    table.cell(3,0).text = "Final Allocation on AY 19-20"

    # Setting Table Data
    table.cell(0,0).text = department + " (" + str(account_num) + ")"
    table.cell(3,2).text = str(s_5hr)
    table.cell(3,3).text = str(s_10hr)
    table.cell(3,4).text = str(p_10hr)
    table.cell(3,5).text = str(p_12hr)
    table.cell(3,6).text = str(p_15hr)
    table.cell(3,7).text = str(total_primary)
    table.cell(3,8).text = str(thanksgiving_hours)
    table.cell(3,9).text = str(spring_hours)
    table.cell(3,10).text = str(christmas_hours)
    table.cell(3,11).text = str(summer_hours)

    # Coloring Cells
    table.rows[0].cells[0]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="939393"/>'.format(nsdecls('w'))))

    table.rows[1].cells[2]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="51B5D8"/>'.format(nsdecls('w'))))
    table.rows[2].cells[2]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="51B5D8"/>'.format(nsdecls('w'))))
    table.rows[2].cells[3]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="51B5D8"/>'.format(nsdecls('w'))))
    table.rows[3].cells[2]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="51B5D8"/>'.format(nsdecls('w'))))
    table.rows[3].cells[3]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="51B5D8"/>'.format(nsdecls('w'))))

    table.rows[1].cells[4]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="F0B27A"/>'.format(nsdecls('w'))))
    table.rows[2].cells[4]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="F0B27A"/>'.format(nsdecls('w'))))
    table.rows[2].cells[5]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="F0B27A"/>'.format(nsdecls('w'))))
    table.rows[2].cells[6]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="F0B27A"/>'.format(nsdecls('w'))))
    table.rows[3].cells[4]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="F0B27A"/>'.format(nsdecls('w'))))
    table.rows[3].cells[5]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="F0B27A"/>'.format(nsdecls('w'))))
    table.rows[3].cells[6]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="F0B27A"/>'.format(nsdecls('w'))))

    table.rows[2].cells[7]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="52BE80"/>'.format(nsdecls('w'))))
    table.rows[3].cells[7]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="52BE80"/>'.format(nsdecls('w'))))

    table.rows[3].cells[0]._tc.get_or_add_tcPr().append(parse_xml(r'<w:shd {} w:fill="#F7DC6F"/>'.format(nsdecls('w'))))


    document.save('%sAllocation-AY1920-%s.docx' % (department[0:8],i))

def extract_data(dataframe,i):
    department = dataframe['Department'][i]
    account_num = dataframe['Account Number'][i]
    try:
        s_5hr = int(dataframe['5hr S'][i])
    except ValueError:
        s_5hr = 0
    try:
        s_10hr = int(dataframe['10hr S'][i])
    except ValueError:
        s_10hr = 0
    try:
        p_10hr = int(dataframe['10hr P'][i])
    except ValueError:
        p_10hr = 0
    try:
        p_12hr = int(dataframe['12hr P'][i])
    except ValueError:
        p_12hr = 0
    try:
        p_15hr = int(dataframe['15hr P'][i])
    except ValueError:
        p_15hr = 0
    try:
        total_primary = int(dataframe['Total Primary Positions'][i])
    except ValueError:
        total_primary = 0
    try:
        thanksgiving_hours = int(dataframe['ThG BRK'][i])
    except ValueError:
        thanksgiving_hours = 0
    try:
        spring_hours = int(dataframe['SP BRK'][i])
    except ValueError:
        spring_hours = 0
    try:
        christmas_hours = int(dataframe['Xmas BRK'][i])
    except ValueError:
        christmas_hours = 0
    try:
        summer_hours = int(dataframe['Summer Hours'][i])
    except ValueError:
        summer_hours = 0

    return department,account_num,s_5hr,s_10hr,p_10hr,p_12hr,p_15hr,total_primary,thanksgiving_hours,spring_hours,christmas_hours,summer_hours

def main():
    df = pandas.read_excel("allocations.xlsx")
    for i in range(len(df['Account Number'])):
        department, account_num, s_5hr, s_10hr, p_10hr,p_12hr,p_15hr,total_primary,thanksgiving_hours,spring_hours,christmas_hours,summer_hours = extract_data(df,i)
        create_doc(department, account_num, s_5hr, s_10hr, p_10hr,p_12hr,p_15hr,total_primary,thanksgiving_hours,spring_hours,christmas_hours,summer_hours,i)

if __name__=="__main__":
    main()
